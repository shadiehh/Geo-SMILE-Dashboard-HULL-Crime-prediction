"""CatBoost Feature top-10, baseline vs. west addition -- one PPTX slide,
using the SAME packed category-coloured-rectangle cluster style as the
existing XGBoost deck (make_feature_stability_pptx.py's draw_cluster_on_slide),
reused directly rather than re-implemented. Scoped to the single slide item
6 asks for, not the full multi-slide XGBoost deck (title, metrics table,
Geo overview, etc.) -- those exist already for XGBoost and were not
requested to be duplicated here.

Writes outputs/att_stability_catboost/feature_top10_baseline_vs_west.pptx.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from make_feature_stability_pptx import _add_textbox, _set_slide_bg, draw_cluster_on_slide
from make_feature_stability_schematic import category_shade_scales, load_feature_categories

ROOT = Path(__file__).resolve().parent
DEST = ROOT / "outputs" / "att_stability_catboost"

BG = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xB0, 0xB8, 0xC0)


def main() -> None:
    baseline_df = pd.read_csv(DEST / "baseline" / "feature_importance.csv")
    west_df = pd.read_csv(DEST / "west" / "feature_importance.csv")

    categories = load_feature_categories()
    scales = category_shade_scales([baseline_df, west_df], categories)
    shared_vmax = max(float(baseline_df["abs_coefficient"].max()), float(west_df["abs_coefficient"].max()))

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide, BG)
    _add_textbox(slide, Inches(0.4), Inches(0.12), Inches(12.5), Inches(0.4),
                 "CatBoost: Feature top-10, Hull-only baseline vs. west addition (E01013112)",
                 size_pt=20, bold=True, colour=WHITE)

    panel_w = Inches(5.9)
    panel_h = Inches(6.05)
    gap = Inches(0.3)
    left0 = Inches(0.4)
    left1 = left0 + panel_w + gap

    draw_cluster_on_slide(
        slide, baseline_df, rep=0, shared_vmax=shared_vmax, categories=categories, scales=scales,
        origin_left=left0, origin_top=Inches(0.72), panel_w=panel_w, panel_h=panel_h,
        title="Hull-only baseline", label_prefix="baseline_", show_neutrals=False,
    )
    draw_cluster_on_slide(
        slide, west_df, rep=4, shared_vmax=shared_vmax, categories=categories, scales=scales,
        origin_left=left1, origin_top=Inches(0.72), panel_w=panel_w, panel_h=panel_h,
        title="West addition", label_prefix="west_", show_neutrals=False,
    )

    _add_textbox(slide, Inches(0.4), Inches(7.08), Inches(12.5), Inches(0.32),
                 "Each rectangle is an editable PowerPoint shape. Size = |SMILE coefficient|, shared scale across "
                 "both panels. Number is rank; value underneath is the coefficient. Top-10 set and order are "
                 "identical between panels (Jaccard = 1.000) -- one real LSOA added to the west of the SMILE "
                 "reference does not change CatBoost's Feature ranking. Not causal.",
                 size_pt=11, colour=MUTED, align=PP_ALIGN.LEFT)

    out_path = DEST / "feature_top10_baseline_vs_west.pptx"
    prs.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
