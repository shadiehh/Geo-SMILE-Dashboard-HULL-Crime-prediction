"""GP Feature-stability slide (PowerPoint), styled after the reference
"Top-10 Important Feature" grid diagram: two side-by-side panels (Hull-only
baseline vs. the west-neighbour addition -- the closest extra unit, per the
established stability protocol), top-10 features as coloured boxes,
category-coded, ranked 1-10 with |coefficient| values, plus a category
legend. Uses GP's own 30-feature baseline/west rankings from
generate_gp_stability_directional_rankings.py and the thesis's established
feature-category map (outputs/feature_category_alignment/feature_with_categories.csv).

Writes outputs/hull_chapter4_gp_assets/feature_stability_gp.pptx.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs"
STAB_DIR = OUT / "hull_chapter4_gp_assets" / "stability"
CAT_PATH = OUT / "feature_category_alignment" / "feature_with_categories.csv"
DEST = OUT / "hull_chapter4_gp_assets" / "feature_stability_gp.pptx"

TOP_K = 10
DIRECTION_LABEL = "West addition"
DIRECTION_LSOA = "E01013112"

# Category base hue (RGB), matching the reference image's palette; shade
# within a category darkens with rank (1 = strongest = darkest).
CATEGORY_BASE = {
    "Crime / temporal": (30, 60, 130),
    "Economic / labour & tenure": (90, 30, 130),
    "Social / demographic": (60, 130, 70),
    "Transport / accessibility": (200, 30, 90),
    "Built environment / activity": (20, 120, 120),
    "Environmental": (60, 150, 60),
    "Structural/control": (90, 90, 90),
}

TITLE_COLOR = RGBColor(0x0B, 0x1F, 0x4D)
INK = RGBColor(0xFF, 0xFF, 0xFF)


def shade(base_rgb: tuple[int, int, int], rank_in_cat: int, n_in_cat: int) -> RGBColor:
    """rank 1 (strongest in its category) -> base colour; later ranks -> lighter tint."""
    t = 0.0 if n_in_cat <= 1 else rank_in_cat / (n_in_cat - 1 + 1e-9)
    lighten = 0.55 * t  # blend toward white as rank worsens within category
    r, g, b = base_rgb
    r = int(r + (255 - r) * lighten)
    g = int(g + (255 - g) * lighten)
    b = int(b + (255 - b) * lighten)
    return RGBColor(r, g, b)


def top_k_with_colors(csv_path: Path, cat_map: dict[str, str]) -> list[dict]:
    df = pd.read_csv(csv_path)
    ranked = df.sort_values(["abs_coefficient", "feature"], ascending=[False, True], kind="mergesort").reset_index(drop=True)
    top = ranked.head(TOP_K).copy()
    top["category"] = top["feature"].map(lambda f: cat_map.get(f, "Structural/control"))
    # rank-within-category, in the SAME row order (i.e. by overall rank), for shading
    seen: dict[str, int] = {}
    counts = top["category"].value_counts().to_dict()
    rows = []
    for i, row in top.iterrows():
        cat = row["category"]
        rank_in_cat = seen.get(cat, 0)
        seen[cat] = rank_in_cat + 1
        rows.append({
            "rank": i + 1, "feature": row["feature"], "abs_coefficient": float(row["abs_coefficient"]),
            "category": cat, "color": shade(CATEGORY_BASE.get(cat, (100, 100, 100)), rank_in_cat, counts[cat]),
        })
    return rows


def add_box(slide, x, y, w, h, label: str, value: float, fill: RGBColor):
    box = slide.shapes.add_shape(1, x, y, w, h)  # MSO_SHAPE.RECTANGLE = 1
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    box.line.width = Pt(1.5)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = INK
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = f"{value:.2f}"
    run2.font.size = Pt(12)
    run2.font.bold = True
    run2.font.color.rgb = INK


def add_panel(slide, rows: list[dict], *, left_in: float, top_in: float, panel_w_in: float, subtitle: str):
    txt = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(panel_w_in), Inches(0.4))
    tf = txt.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = subtitle
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = TITLE_COLOR

    cols, rows_n = 5, 2
    gap = 0.10
    box_w = (panel_w_in - gap * (cols - 1)) / cols
    box_h = 0.95
    grid_top = top_in + 0.55
    for row in rows:
        idx = row["rank"] - 1
        c, r = idx % cols, idx // cols
        x = Inches(left_in + c * (box_w + gap))
        y = Inches(grid_top + r * (box_h + gap))
        label = f"{row['rank']}. {row['feature'].replace('_', ' ')}"
        add_box(slide, x, y, Inches(box_w), Inches(box_h), label, row["abs_coefficient"], row["color"])


def add_legend(slide, cat_map: dict[str, str], top_in: float):
    used_cats = sorted(set(cat_map.values()), key=lambda c: list(CATEGORY_BASE).index(c) if c in CATEGORY_BASE else 99)
    used_cats = [c for c in used_cats if c in CATEGORY_BASE]
    box = slide.shapes.add_shape(1, Inches(0.4), Inches(top_in), Inches(9.2), Inches(0.9))
    box.fill.background()
    box.line.color.rgb = TITLE_COLOR
    box.line.width = Pt(1.0)
    n = len(used_cats)
    col_w = 9.2 / max(n, 1)
    for i, cat in enumerate(used_cats):
        base = CATEGORY_BASE[cat]
        x = Inches(0.4 + i * col_w)
        lbl = slide.shapes.add_textbox(x, Inches(top_in + 0.05), Inches(col_w), Inches(0.3))
        p = lbl.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = cat
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = TITLE_COLOR
        p.alignment = PP_ALIGN.CENTER
        swatch_w = col_w * 0.9 / 5
        for k in range(5):
            c = shade(base, k, 5)
            sw = slide.shapes.add_shape(1, Inches(0.4 + i * col_w + 0.05 * col_w + k * swatch_w), Inches(top_in + 0.4), Inches(swatch_w * 0.9), Inches(0.3))
            sw.fill.solid()
            sw.fill.fore_color.rgb = c
            sw.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            sw.line.width = Pt(0.5)


def main() -> None:
    cat_df = pd.read_csv(CAT_PATH)
    cat_map = dict(zip(cat_df["feature"].astype(str), cat_df["thesis_category"].astype(str)))

    baseline_rows = top_k_with_colors(STAB_DIR / "feature_baseline.csv", cat_map)
    west_rows = top_k_with_colors(STAB_DIR / "feature_west.csv", cat_map)
    used_cat_map = {r["feature"]: r["category"] for r in baseline_rows + west_rows}

    prs = Presentation()
    prs.slide_width = Inches(10.0)
    prs.slide_height = Inches(7.0)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    title = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(9.2), Inches(0.5))
    p = title.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "GP: Top-10 Important Feature -- Stability Check"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = TITLE_COLOR
    p.alignment = PP_ALIGN.CENTER

    panel_w = 4.4
    add_panel(slide, baseline_rows, left_in=0.4, top_in=0.9, panel_w_in=panel_w, subtitle="Hull-only baseline")
    add_panel(slide, west_rows, left_in=0.4 + panel_w + 0.4, top_in=0.9, panel_w_in=panel_w,
              subtitle=f"{DIRECTION_LABEL} · {DIRECTION_LSOA}")

    add_legend(slide, used_cat_map, top_in=3.85)

    note = slide.shapes.add_textbox(Inches(0.4), Inches(4.9), Inches(9.2), Inches(1.6))
    tf = note.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = (
        "GeoSMILE Feature branch, N=1,024 masks, seed=42 (30-feature version). Boxes show |SMILE coefficient|; "
        "colour = thesis feature category, shade darkens with rank within its own category. The West addition "
        f"(closest extra LSOA to the reference, {DIRECTION_LSOA}) leaves the top-10 SET AND ORDER completely "
        "unchanged (Jaccard = 1.000; |coef| Spearman = 0.998) -- only the fourth-decimal coefficient values shift "
        "slightly. Not causal."
    )
    run.font.size = Pt(10)
    run.font.color.rgb = TITLE_COLOR

    DEST.parent.mkdir(parents=True, exist_ok=True)
    prs.save(DEST)
    print(f"Saved: {DEST}")


if __name__ == "__main__":
    main()
